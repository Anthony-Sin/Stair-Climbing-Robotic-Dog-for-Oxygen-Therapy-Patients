#!/usr/bin/env python3
"""Render the ACTUAL generated .pptx to PNGs by reading each shape's geometry,
text and images back out with python-pptx and drawing them with Pillow.
Font substitution (DejaVu/Liberation) differs from Segoe UI but layout,
overflow, overlaps, colors and content are faithful."""
import os, io, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

PPTX = sys.argv[1] if len(sys.argv)>1 else "TASH_pitch.pptx"
OUT = "preview"; os.makedirs(OUT, exist_ok=True)
DPI = 120
prs = Presentation(PPTX)
SW = int(prs.slide_width/914400*DPI); SH = int(prs.slide_height/914400*DPI)

FONTS = {
 ("disp",False):"/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
 ("disp",True):"/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
 ("mono",False):"/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
 ("mono",True):"/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf",
}
_cache={}
def font(kind,bold,px):
    k=(kind,bold,px)
    if k not in _cache:
        _cache[k]=ImageFont.truetype(FONTS[(kind,bold)], max(6,int(px)))
    return _cache[k]

def emu_px(v): return int(v/914400*DPI)
def rgb(c):
    try: return (c[0],c[1],c[2])
    except Exception: return None

def wrap(draw, runs, fnt_for, maxw):
    """runs=[(text,bold)]; returns list of lines, each line=[(word,bold)]."""
    lines=[[]]; x=0
    for text,bold in runs:
        for tok in text.replace("\n"," \n ").split(" "):
            if tok=="\n": lines.append([]); x=0; continue
            if tok=="": continue
            f=fnt_for(bold); w=draw.textlength(tok+" ", font=f)
            if x+w>maxw and lines[-1]:
                lines.append([]); x=0
            lines[-1].append((tok,bold)); x+=w
    return lines

for si, slide in enumerate(prs.slides):
    img=Image.new("RGB",(SW,SH),(27,27,24)); d=ImageDraw.Draw(img)
    for sh in slide.shapes:
        try: L,T,W,H=emu_px(sh.left),emu_px(sh.top),emu_px(sh.width),emu_px(sh.height)
        except Exception: continue
        st=sh.shape_type
        # picture / movie poster
        if st==MSO_SHAPE_TYPE.PICTURE or st==MSO_SHAPE_TYPE.MEDIA:
            try:
                blob=sh.image.blob; im=Image.open(io.BytesIO(blob)).convert("RGB")
                im=im.resize((max(1,W),max(1,H))); img.paste(im,(L,T))
                if st==MSO_SHAPE_TYPE.MEDIA:
                    d.rectangle([L,T,L+W,H+T],outline=(192,125,60),width=2)
                continue
            except Exception: pass
        # auto shapes: fill + border
        fillc=None; borderc=None
        try:
            if sh.fill.type==1: fillc=rgb(sh.fill.fore_color.rgb)
        except Exception: pass
        try:
            if sh.line.width and sh.line.width>0: borderc=rgb(sh.line.color.rgb)
        except Exception: pass
        if st==MSO_SHAPE_TYPE.AUTO_SHAPE or fillc or borderc:
            if fillc: d.rounded_rectangle([L,T,L+W,T+H],radius=6,fill=fillc)
            if borderc: d.rounded_rectangle([L,T,L+W,T+H],radius=6,outline=borderc,width=2)
        # text
        if sh.has_text_frame:
            y=T+4
            for para in sh.text_frame.paragraphs:
                runs=[(r.text, bool(r.font.bold)) for r in para.runs if r.text]
                if not runs:
                    y+=10; continue
                sz=None; kind="disp"; col=(47,44,40)
                for r in para.runs:
                    if r.font.size: sz=r.font.size.pt
                    if r.font.name and "Consol" in (r.font.name or ""): kind="mono"
                    if r.font.name and "Mono" in (r.font.name or ""): kind="mono"
                    try:
                        if r.font.color and r.font.color.rgb: col=rgb(r.font.color.rgb)
                    except Exception: pass
                sz=sz or 12
                px=sz*DPI/72.0
                fnt_for=lambda b: font(kind,b,px)
                align=str(para.alignment) if para.alignment else "LEFT"
                lines=wrap(d,runs,fnt_for,W-8)
                for line in lines:
                    lx=L+4
                    lw=sum(d.textlength(w+" ",font=fnt_for(b)) for w,b in line)
                    if "CENTER" in align: lx=L+(W-lw)/2
                    for word,bold in line:
                        f=fnt_for(bold)
                        d.text((lx,y),word,font=f,fill=col)
                        lx+=d.textlength(word+" ",font=f)
                    y+=px*1.28
    img.save(f"{OUT}/slide_{si:02d}.png")
print("rendered", len(prs.slides._sldIdLst), "slides to", OUT)
